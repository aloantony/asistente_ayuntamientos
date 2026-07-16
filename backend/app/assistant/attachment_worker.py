"""Resource-limited parser process for untrusted assistant attachments.

This module is executed with ``python -I``.  The API process deliberately does
not import PDF or spreadsheet parser libraries; potentially hostile structured
documents are handled only after Linux resource limits and archive checks have
been applied here.
"""

from __future__ import annotations

import argparse
import html
import json
import re
import resource
import stat
import sys
import zipfile
from pathlib import Path, PurePosixPath

MAX_PDF_PAGES = 20
MAX_DOCX_TEXT_FRAGMENTS = 2_000
MAX_XLSX_SHEETS = 10
MAX_XLSX_ROWS_PER_SHEET = 200
MAX_XLSX_CELLS_PER_ROW = 50
_DOCX_TEXT_RE = re.compile(rb"<w:t(?:\s[^>]*)?>(.*?)</w:t>", re.DOTALL)


class RejectedAttachment(Exception):
    def __init__(self, status: str, reason: str) -> None:
        super().__init__(reason)
        self.status = status
        self.reason = reason


def _set_limit(limit: int, value: int) -> None:
    current_soft, current_hard = resource.getrlimit(limit)
    hard = value if current_hard == resource.RLIM_INFINITY else min(value, current_hard)
    soft = min(value, hard)
    resource.setrlimit(limit, (soft, hard))


def apply_resource_limits(args: argparse.Namespace) -> None:
    if sys.platform != "linux":
        raise RejectedAttachment("failed", "isolated parsing requires Linux")
    _set_limit(resource.RLIMIT_CPU, args.cpu_seconds)
    _set_limit(resource.RLIMIT_AS, args.memory_bytes)
    _set_limit(resource.RLIMIT_NOFILE, args.max_fds)
    _set_limit(resource.RLIMIT_NPROC, 1)
    _set_limit(resource.RLIMIT_CORE, 0)
    _set_limit(resource.RLIMIT_FSIZE, 1024 * 1024)


def _read_magic(path: Path, length: int = 8) -> bytes:
    with path.open("rb") as source:
        return source.read(length)


def _validate_pdf_magic(path: Path) -> None:
    if not _read_magic(path).startswith(b"%PDF-"):
        raise RejectedAttachment("failed", "invalid PDF signature")


def _safe_archive_member_name(name: str) -> str:
    if not name or "\x00" in name or "\\" in name or "//" in name:
        raise RejectedAttachment("failed", "unsafe archive member name")
    path = PurePosixPath(name)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise RejectedAttachment("failed", "unsafe archive member path")
    return path.as_posix().rstrip("/")


def validate_zip_archive(
    path: Path,
    *,
    max_members: int,
    max_member_bytes: int,
    max_total_bytes: int,
    max_compression_ratio: float,
    required_member: str,
) -> None:
    if not _read_magic(path).startswith(b"PK\x03\x04"):
        raise RejectedAttachment("failed", "invalid ZIP signature")

    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > max_members:
                raise RejectedAttachment("too_large", "too many archive members")

            total_size = 0
            seen_names: set[str] = set()
            safe_names: set[str] = set()
            for info in infos:
                safe_name = _safe_archive_member_name(info.filename)
                folded_name = safe_name.casefold()
                if folded_name in seen_names:
                    raise RejectedAttachment("failed", "duplicate archive member")
                seen_names.add(folded_name)
                safe_names.add(safe_name)

                if info.flag_bits & 0x1:
                    raise RejectedAttachment("failed", "encrypted archive member")
                unix_mode = info.external_attr >> 16
                if unix_mode and stat.S_ISLNK(unix_mode):
                    raise RejectedAttachment("failed", "archive symlink")
                if info.is_dir():
                    continue
                if info.file_size > max_member_bytes:
                    raise RejectedAttachment("too_large", "archive member too large")
                total_size += info.file_size
                if total_size > max_total_bytes:
                    raise RejectedAttachment("too_large", "archive expands too much")
                if info.file_size:
                    ratio = info.file_size / max(1, info.compress_size)
                    if ratio > max_compression_ratio:
                        raise RejectedAttachment(
                            "too_large",
                            "archive compression ratio too high",
                        )

            if required_member not in safe_names:
                raise RejectedAttachment("failed", "required archive member missing")
            bad_member = archive.testzip()
            if bad_member is not None:
                raise RejectedAttachment("failed", "corrupt archive member")
    except RejectedAttachment:
        raise
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as error:
        raise RejectedAttachment("failed", "invalid ZIP archive") from error


def extract_pdf(path: Path, max_chars: int) -> str:
    _validate_pdf_magic(path)
    # Imported only after rlimits are active in this isolated process.
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=True)
    if reader.is_encrypted:
        raise RejectedAttachment("failed", "encrypted PDF")
    parts: list[str] = []
    chars = 0
    for page_number, page in enumerate(reader.pages):
        if page_number >= MAX_PDF_PAGES:
            break
        page_text = page.extract_text() or ""
        remaining = max_chars - chars
        if remaining <= 0:
            break
        if page_text:
            chunk = page_text[:remaining]
            parts.append(chunk)
            chars += len(chunk)
    return "\n\n".join(parts)


def extract_docx(path: Path, max_chars: int) -> str:
    with zipfile.ZipFile(path) as archive:
        raw_xml = archive.read("word/document.xml")
    fragments = [
        html.unescape(match.decode("utf-8", errors="replace"))
        for match in _DOCX_TEXT_RE.findall(raw_xml)[:MAX_DOCX_TEXT_FRAGMENTS]
    ]
    return "\n".join(fragments)[:max_chars]


def extract_xlsx(path: Path, max_chars: int) -> str:
    # Imported only after rlimits and ZIP policy checks are active.
    from openpyxl import load_workbook

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    parts: list[str] = []
    chars = 0
    try:
        for worksheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            header = f"Hoja: {worksheet.title}"
            parts.append(header)
            chars += len(header) + 1
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True)):
                if row_number >= MAX_XLSX_ROWS_PER_SHEET:
                    break
                line = " | ".join(
                    str(value).strip()
                    for value in row[:MAX_XLSX_CELLS_PER_ROW]
                    if value is not None and str(value).strip()
                )
                if not line:
                    continue
                remaining = max_chars - chars
                if remaining <= 0:
                    return "\n".join(parts)
                chunk = line[:remaining]
                parts.append(chunk)
                chars += len(chunk) + 1
    finally:
        workbook.close()
    return "\n".join(parts)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--path", required=True)
    parser.add_argument("--content-type", required=True)
    parser.add_argument("--max-chars", type=int, required=True)
    parser.add_argument("--max-total-bytes", type=int, required=True)
    parser.add_argument("--max-member-bytes", type=int, required=True)
    parser.add_argument("--max-members", type=int, required=True)
    parser.add_argument("--max-compression-ratio", type=float, required=True)
    parser.add_argument("--cpu-seconds", type=int, required=True)
    parser.add_argument("--memory-bytes", type=int, required=True)
    parser.add_argument("--max-fds", type=int, required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        apply_resource_limits(args)
        path = Path(args.path)
        if args.content_type == "application/pdf":
            text = extract_pdf(path, args.max_chars)
        elif args.content_type == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            validate_zip_archive(
                path,
                max_members=args.max_members,
                max_member_bytes=args.max_member_bytes,
                max_total_bytes=args.max_total_bytes,
                max_compression_ratio=args.max_compression_ratio,
                required_member="word/document.xml",
            )
            text = extract_docx(path, args.max_chars)
        elif args.content_type == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ):
            validate_zip_archive(
                path,
                max_members=args.max_members,
                max_member_bytes=args.max_member_bytes,
                max_total_bytes=args.max_total_bytes,
                max_compression_ratio=args.max_compression_ratio,
                required_member="[Content_Types].xml",
            )
            text = extract_xlsx(path, args.max_chars)
        else:
            raise RejectedAttachment("failed", "unsupported parser content type")
        payload = {"ok": True, "status": "ready", "text": text[: args.max_chars]}
    except RejectedAttachment as error:
        payload = {"ok": False, "status": error.status, "error": error.reason}
    except (MemoryError, OSError, RuntimeError, ValueError, zipfile.BadZipFile):
        payload = {"ok": False, "status": "failed", "error": "parser failed"}
    except Exception:
        payload = {"ok": False, "status": "failed", "error": "parser failed"}

    sys.stdout.write(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
