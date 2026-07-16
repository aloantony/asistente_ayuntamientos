"""Synchronize the official Castilla y Leon municipality directory.

The directory source is the current INE municipality catalogue, while
population figures come from the separately reviewed official population
snapshot. The command is a dry-run unless ``--apply`` is supplied.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.db.model_registry import register_all_models
from app.db.session import SessionLocal
from app.municipalities.ine_population import (
    DOWNLOAD_CHUNK_BYTES,
    DOWNLOAD_TIMEOUT_SECONDS,
    INE_SOURCE_SPECS,
    MAX_WORKBOOK_BYTES,
    InePopulationError,
    InePopulationRecord,
    PopulationProvenance,
    _density,
    _normalized_name,
    _population_conflict,
    download_ine_archive,
    load_ine_population,
    read_local_archive,
)
from app.municipalities.models import Municipality

CASTILLA_Y_LEON_CODE = "07"
CASTILLA_Y_LEON_NAME = "Castilla y León"
COUNTRY_NAME = "España"
PROVINCES = {
    "05": "Ávila",
    "09": "Burgos",
    "24": "León",
    "34": "Palencia",
    "37": "Salamanca",
    "40": "Segovia",
    "42": "Soria",
    "47": "Valladolid",
    "49": "Zamora",
}
EXPECTED_PROVINCE_COUNTS = {
    "05": 248,
    "09": 371,
    "24": 211,
    "34": 191,
    "37": 362,
    "40": 209,
    "42": 183,
    "47": 225,
    "49": 248,
}


class IneDirectoryError(ValueError):
    """The official directory or synchronization plan is not safe to apply."""


@dataclass(frozen=True)
class IneDirectorySourceSpec:
    year: int
    reference_date: date
    workbook_url: str
    title: str
    headers: tuple[str, ...]
    expected_sha256: str
    expected_rows: int
    expected_region_rows: int
    expected_province_counts: tuple[tuple[str, int], ...]


INE_DIRECTORY_SOURCE_SPECS = {
    2026: IneDirectorySourceSpec(
        year=2026,
        reference_date=date(2026, 1, 1),
        workbook_url=(
            "https://www.ine.es/daco/daco42/codmun/diccionario26.xlsx"
        ),
        title=(
            "Relación de municipios y códigos por comunidades autónomas "
            "y provincias a 1 de enero de 2026"
        ),
        headers=("CODAUTO", "CPRO", "CMUN", "DC", "NOMBRE"),
        expected_sha256=(
            "07f8e8d64eba73fe9d425196fe82f4e09a9886a287abd650e88ca8d98dd49052"
        ),
        expected_rows=8132,
        expected_region_rows=2248,
        expected_province_counts=tuple(EXPECTED_PROVINCE_COUNTS.items()),
    )
}


@dataclass(frozen=True)
class IneDirectoryRecord:
    autonomous_community_code: str
    province_code: str
    municipality_code: str
    check_digit: str
    name: str

    @property
    def ine_code(self) -> str:
        return f"{self.province_code}{self.municipality_code}"

    @property
    def province(self) -> str:
        return PROVINCES[self.province_code]


@dataclass(frozen=True)
class DirectoryProvenance:
    reference_date: date
    source_url: str
    source_sha256: str


@dataclass(frozen=True)
class DirectoryIssue:
    reason: str
    detail: str
    ine_code: str | None = None
    municipality_name: str | None = None
    municipality_id: int | None = None


@dataclass
class DirectoryMutation:
    municipality: Municipality | None
    record: IneDirectoryRecord
    population: InePopulationRecord
    values: dict[str, object]


@dataclass
class DirectorySyncPlan:
    directory_provenance: DirectoryProvenance
    population_provenance: PopulationProvenance
    source_rows: int
    source_by_province: dict[str, int]
    source_population: int
    source_men: int
    source_women: int
    database_rows: int
    official_rows_before: int
    unchanged_rows: int = 0
    creates: list[DirectoryMutation] = field(default_factory=list)
    updates: list[DirectoryMutation] = field(default_factory=list)
    missing_population_codes: list[DirectoryIssue] = field(default_factory=list)
    population_name_mismatches: list[DirectoryIssue] = field(default_factory=list)
    population_conflicts: list[DirectoryIssue] = field(default_factory=list)
    duplicate_database_codes: list[DirectoryIssue] = field(default_factory=list)
    legacy_match_conflicts: list[DirectoryIssue] = field(default_factory=list)
    extra_region_rows: list[DirectoryIssue] = field(default_factory=list)
    official_name_changes: list[DirectoryIssue] = field(default_factory=list)

    @property
    def blocking_issues(self) -> list[DirectoryIssue]:
        return [
            *self.missing_population_codes,
            *self.population_conflicts,
            *self.duplicate_database_codes,
            *self.legacy_match_conflicts,
        ]

    def summary(self, *, applied: bool) -> dict[str, Any]:
        return {
            "ok": not self.blocking_issues,
            "mode": "apply" if applied else "dry-run",
            "directory_source": {
                "reference_date": self.directory_provenance.reference_date.isoformat(),
                "url": self.directory_provenance.source_url,
                "sha256": self.directory_provenance.source_sha256,
                "rows": self.source_rows,
                "by_province": self.source_by_province,
            },
            "population_source": {
                "reference_year": self.population_provenance.reference_year,
                "url": self.population_provenance.source_url,
                "sha256": self.population_provenance.source_sha256,
                "total": self.source_population,
                "men": self.source_men,
                "women": self.source_women,
                "stored_fields": ["population"],
                "available_not_stored": ["men", "women"],
            },
            "database": {
                "rows_before": self.database_rows,
                "official_rows_before": self.official_rows_before,
                "would_create": len(self.creates),
                "would_update": len(self.updates),
                "created": len(self.creates) if applied else 0,
                "updated": len(self.updates) if applied else 0,
                "unchanged": self.unchanged_rows,
                "official_rows_after": (
                    self.official_rows_before + len(self.creates)
                    if applied
                    else self.official_rows_before
                ),
            },
            "issues": {
                "missing_population_codes": [
                    asdict(issue) for issue in self.missing_population_codes
                ],
                "population_name_mismatches": [
                    asdict(issue) for issue in self.population_name_mismatches
                ],
                "population_conflicts": [
                    asdict(issue) for issue in self.population_conflicts
                ],
                "duplicate_database_codes": [
                    asdict(issue) for issue in self.duplicate_database_codes
                ],
                "legacy_match_conflicts": [
                    asdict(issue) for issue in self.legacy_match_conflicts
                ],
                "extra_region_rows": [
                    asdict(issue) for issue in self.extra_region_rows
                ],
                "official_name_changes": [
                    asdict(issue) for issue in self.official_name_changes
                ],
            },
        }


def download_directory_workbook(
    spec: IneDirectorySourceSpec,
    *,
    max_bytes: int = MAX_WORKBOOK_BYTES,
    timeout_seconds: float = DOWNLOAD_TIMEOUT_SECONDS,
) -> bytes:
    request = Request(
        spec.workbook_url,
        headers={"User-Agent": "asistente-ayuntamientos-ine-directory/1.0"},
    )
    try:
        with urlopen(request, timeout=timeout_seconds) as response:
            final_url = urlsplit(response.geturl())
            expected_url = urlsplit(spec.workbook_url)
            if (
                final_url.scheme != "https"
                or final_url.hostname != expected_url.hostname
            ):
                raise IneDirectoryError(
                    "La descarga del directorio INE redirigió a un origen no permitido"
                )

            content_length = response.headers.get("Content-Length")
            if content_length is not None and int(content_length) > max_bytes:
                raise IneDirectoryError(
                    "El XLSX del directorio INE supera el tamaño permitido"
                )

            workbook = bytearray()
            while chunk := response.read(DOWNLOAD_CHUNK_BYTES):
                workbook.extend(chunk)
                if len(workbook) > max_bytes:
                    raise IneDirectoryError(
                        "El XLSX del directorio INE supera el tamaño permitido"
                    )
    except IneDirectoryError:
        raise
    except (HTTPError, URLError, OSError, ValueError) as exc:
        raise IneDirectoryError(
            f"No se pudo descargar el directorio INE: {exc}"
        ) from exc

    if not workbook:
        raise IneDirectoryError("El XLSX descargado del directorio INE está vacío")
    return bytes(workbook)


def read_local_workbook(
    path: Path,
    *,
    max_bytes: int = MAX_WORKBOOK_BYTES,
) -> bytes:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise IneDirectoryError(f"No se pudo leer el XLSX local: {exc}") from exc
    if size <= 0:
        raise IneDirectoryError("El XLSX local está vacío")
    if size > max_bytes:
        raise IneDirectoryError("El XLSX local supera el tamaño permitido")
    try:
        return path.read_bytes()
    except OSError as exc:
        raise IneDirectoryError(f"No se pudo leer el XLSX local: {exc}") from exc


def load_castilla_y_leon_directory(
    workbook_bytes: bytes,
    spec: IneDirectorySourceSpec,
) -> tuple[dict[str, IneDirectoryRecord], DirectoryProvenance]:
    received_sha256 = hashlib.sha256(workbook_bytes).hexdigest()
    if received_sha256 != spec.expected_sha256:
        raise IneDirectoryError(
            "El directorio INE no coincide con la versión revisada; "
            f"esperado {spec.expected_sha256}, recibido {received_sha256}"
        )

    try:
        workbook = load_workbook(
            filename=io.BytesIO(workbook_bytes),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise IneDirectoryError(
            f"No se pudo abrir el directorio INE: {exc}"
        ) from exc

    try:
        if len(workbook.worksheets) != 1:
            raise IneDirectoryError("El directorio INE debe contener una sola hoja")
        rows = iter(workbook.worksheets[0].iter_rows(values_only=True))
        try:
            title_row = next(rows)
            header_row = next(rows)
        except StopIteration as exc:
            raise IneDirectoryError(
                "El directorio INE no contiene cabeceras"
            ) from exc

        title = _required_text(title_row[0] if title_row else None, "título", 1)
        if title != spec.title:
            raise IneDirectoryError(
                f"Título inesperado en el directorio INE: {title!r}"
            )
        headers = tuple(
            _required_text(value, f"cabecera {index + 1}", 2)
            for index, value in enumerate(header_row[: len(spec.headers)])
        )
        if headers != spec.headers:
            raise IneDirectoryError(
                f"Cabeceras inesperadas en el directorio INE: {headers!r}"
            )

        all_records: dict[str, IneDirectoryRecord] = {}
        region_records: dict[str, IneDirectoryRecord] = {}
        for row_number, row in enumerate(rows, start=3):
            values = tuple(row[: len(spec.headers)])
            if not values or all(value is None or value == "" for value in values):
                continue
            if len(values) != len(spec.headers):
                raise IneDirectoryError(
                    f"Fila {row_number}: se esperaban {len(spec.headers)} columnas"
                )

            record = IneDirectoryRecord(
                autonomous_community_code=_code_component(
                    values[0], 2, "CODAUTO", row_number
                ),
                province_code=_code_component(values[1], 2, "CPRO", row_number),
                municipality_code=_code_component(
                    values[2], 3, "CMUN", row_number
                ),
                check_digit=_code_component(values[3], 1, "DC", row_number),
                name=_required_text(values[4], "NOMBRE", row_number),
            )
            if record.ine_code in all_records:
                raise IneDirectoryError(
                    f"Fila {row_number}: código INE duplicado {record.ine_code}"
                )
            all_records[record.ine_code] = record

            if record.autonomous_community_code != CASTILLA_Y_LEON_CODE:
                continue
            if record.province_code not in PROVINCES:
                raise IneDirectoryError(
                    f"Fila {row_number}: provincia inesperada para Castilla y León"
                )
            region_records[record.ine_code] = record
    finally:
        workbook.close()

    if len(all_records) != spec.expected_rows:
        raise IneDirectoryError(
            "Número inesperado de municipios en el directorio INE: "
            f"esperado {spec.expected_rows}, recibido {len(all_records)}"
        )
    if len(region_records) != spec.expected_region_rows:
        raise IneDirectoryError(
            "Número inesperado de municipios de Castilla y León: "
            f"esperado {spec.expected_region_rows}, recibido {len(region_records)}"
        )

    received_counts = Counter(
        record.province_code for record in region_records.values()
    )
    expected_counts = dict(spec.expected_province_counts)
    if dict(received_counts) != expected_counts:
        raise IneDirectoryError(
            "Recuento provincial inesperado en el directorio INE: "
            f"esperado {expected_counts}, recibido {dict(received_counts)}"
        )

    return region_records, DirectoryProvenance(
        reference_date=spec.reference_date,
        source_url=spec.workbook_url,
        source_sha256=received_sha256,
    )


def build_directory_sync_plan(
    db: Session,
    directory_records: dict[str, IneDirectoryRecord],
    population_records: dict[str, InePopulationRecord],
    directory_provenance: DirectoryProvenance,
    population_provenance: PopulationProvenance,
    *,
    overwrite_existing: bool = False,
) -> DirectorySyncPlan:
    register_all_models()
    municipalities = list(
        db.scalars(select(Municipality).order_by(Municipality.id)).all()
    )
    official_codes = set(directory_records)
    database_by_code: dict[str, Municipality] = {}
    legacy_by_identity: dict[tuple[str, str], list[Municipality]] = {}
    legacy_by_candidate_code: dict[str, list[Municipality]] = {}
    duplicate_database_codes: list[DirectoryIssue] = []
    for municipality in municipalities:
        code = (municipality.ine_code or "").strip()
        if not code:
            identity = (
                _normalized_name(municipality.province),
                _normalized_name(municipality.name),
            )
            legacy_by_identity.setdefault(identity, []).append(municipality)
            continue
        if code not in official_codes:
            identity = (
                _normalized_name(municipality.province),
                _normalized_name(municipality.name),
            )
            legacy_by_identity.setdefault(identity, []).append(municipality)
            candidate_code = _legacy_official_code_candidate(code, official_codes)
            if candidate_code is not None:
                legacy_by_candidate_code.setdefault(candidate_code, []).append(
                    municipality
                )
        existing = database_by_code.get(code)
        if existing is not None:
            duplicate_database_codes.append(
                DirectoryIssue(
                    reason="duplicate_database_code",
                    detail=f"También figura en el municipio {existing.id}",
                    ine_code=code,
                    municipality_name=municipality.name,
                    municipality_id=municipality.id,
                )
            )
            continue
        database_by_code[code] = municipality

    selected_populations = [
        population_records[code]
        for code in official_codes
        if code in population_records
    ]
    plan = DirectorySyncPlan(
        directory_provenance=directory_provenance,
        population_provenance=population_provenance,
        source_rows=len(directory_records),
        source_by_province={
            PROVINCES[code]: count
            for code, count in Counter(
                record.province_code for record in directory_records.values()
            ).items()
        },
        source_population=sum(record.population for record in selected_populations),
        source_men=sum(record.men for record in selected_populations),
        source_women=sum(record.women for record in selected_populations),
        database_rows=len(municipalities),
        official_rows_before=sum(
            code in database_by_code for code in official_codes
        ),
        duplicate_database_codes=duplicate_database_codes,
    )

    for municipality in municipalities:
        code = (municipality.ine_code or "").strip()
        if (
            _normalized_name(municipality.autonomous_community)
            == _normalized_name(CASTILLA_Y_LEON_NAME)
            and code not in official_codes
        ):
            plan.extra_region_rows.append(
                DirectoryIssue(
                    reason="not_in_official_directory",
                    detail="La fila se conserva y requiere revisión manual",
                    ine_code=municipality.ine_code,
                    municipality_name=municipality.name,
                    municipality_id=municipality.id,
                )
            )

    for ine_code, record in directory_records.items():
        population = population_records.get(ine_code)
        if population is None:
            plan.missing_population_codes.append(
                DirectoryIssue(
                    reason="missing_population",
                    detail="No aparece en la fuente de población seleccionada",
                    ine_code=ine_code,
                    municipality_name=record.name,
                )
            )
            continue
        if _normalized_name(population.name) != _normalized_name(record.name):
            plan.population_name_mismatches.append(
                DirectoryIssue(
                    reason="population_name_mismatch",
                    detail=f"Población {population_provenance.reference_year}: {population.name}",
                    ine_code=ine_code,
                    municipality_name=record.name,
                )
            )

        municipality = database_by_code.get(ine_code)
        if municipality is None:
            conflicting_ids: set[int] = set()
            for candidate in legacy_by_candidate_code.get(ine_code, []):
                conflicting_ids.add(candidate.id)
                plan.legacy_match_conflicts.append(
                    DirectoryIssue(
                        reason="possible_legacy_code",
                        detail=(
                            f"El código no canónico {candidate.ine_code!r} puede "
                            f"corresponder a {ine_code}; debe reconciliarse manualmente"
                        ),
                        ine_code=ine_code,
                        municipality_name=record.name,
                        municipality_id=candidate.id,
                    )
                )
            identity = (
                _normalized_name(record.province),
                _normalized_name(record.name),
            )
            legacy_candidates = legacy_by_identity.get(identity, [])
            if legacy_candidates:
                for candidate in legacy_candidates:
                    if candidate.id in conflicting_ids:
                        continue
                    conflicting_ids.add(candidate.id)
                    plan.legacy_match_conflicts.append(
                        DirectoryIssue(
                            reason="possible_legacy_match",
                            detail=(
                                "Existe una fila sin código con el mismo nombre y "
                                "provincia; debe reconciliarse manualmente"
                            ),
                            ine_code=ine_code,
                            municipality_name=record.name,
                            municipality_id=candidate.id,
                        )
                    )
            if conflicting_ids:
                continue
        if municipality is not None:
            conflict = _population_conflict(
                municipality,
                population,
                population_provenance,
                overwrite_existing=overwrite_existing,
            )
            if conflict is not None:
                plan.population_conflicts.append(
                    DirectoryIssue(
                        reason=conflict.reason,
                        detail=conflict.detail,
                        ine_code=conflict.ine_code,
                        municipality_name=conflict.municipality_name,
                        municipality_id=conflict.municipality_id,
                    )
                )
                continue
            if municipality.name != record.name:
                plan.official_name_changes.append(
                    DirectoryIssue(
                        reason="official_name_change",
                        detail=f"Nombre anterior: {municipality.name}",
                        ine_code=ine_code,
                        municipality_name=record.name,
                        municipality_id=municipality.id,
                    )
                )

        surface_km2 = None if municipality is None else municipality.surface_km2
        values: dict[str, object] = {
            "name": record.name,
            "province": record.province,
            "autonomous_community": CASTILLA_Y_LEON_NAME,
            "country": COUNTRY_NAME,
            "ine_code": ine_code,
            "ine_check_digit": record.check_digit,
            "directory_reference_date": directory_provenance.reference_date,
            "directory_source_url": directory_provenance.source_url,
            "directory_source_sha256": directory_provenance.source_sha256,
            "population": population.population,
            "population_reference_year": population_provenance.reference_year,
            "population_source_url": population_provenance.source_url,
            "population_source_sha256": population_provenance.source_sha256,
            "density": _density(population.population, surface_km2),
            "municipality_type": "municipality",
            "status": "active",
        }
        mutation = DirectoryMutation(
            municipality=municipality,
            record=record,
            population=population,
            values=values,
        )
        if municipality is None:
            plan.creates.append(mutation)
            continue
        if all(getattr(municipality, key) == value for key, value in values.items()):
            plan.unchanged_rows += 1
        else:
            plan.updates.append(mutation)

    return plan


def apply_directory_sync(db: Session, plan: DirectorySyncPlan) -> None:
    if plan.blocking_issues:
        raise IneDirectoryError(
            "La sincronización contiene conflictos; revisa el dry-run antes de aplicar"
        )
    try:
        for mutation in plan.creates:
            db.add(Municipality(**mutation.values))
        for mutation in plan.updates:
            if mutation.municipality is None:
                raise IneDirectoryError("El plan de actualización no es válido")
            for key, value in mutation.values.items():
                setattr(mutation.municipality, key, value)
        db.commit()
    except Exception:
        db.rollback()
        raise


def _required_text(value: object, field_name: str, row_number: int) -> str:
    text = str(value or "").strip()
    if not text:
        raise IneDirectoryError(f"Fila {row_number}: {field_name} está vacío")
    return text


def _code_component(
    value: object,
    width: int,
    field_name: str,
    row_number: int,
) -> str:
    if isinstance(value, bool):
        raise IneDirectoryError(f"Fila {row_number}: {field_name} no es válido")
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value or "").strip()
    if not text.isdigit() or len(text) > width:
        raise IneDirectoryError(f"Fila {row_number}: {field_name} no es válido")
    return text.zfill(width)


def _legacy_official_code_candidate(
    code: str,
    official_codes: set[str],
) -> str | None:
    if not code.isdigit():
        return None
    if len(code) < 5:
        candidate = code.zfill(5)
    elif len(code) == 6:
        candidate = code[:5]
    else:
        return None
    return candidate if candidate in official_codes else None


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Sincroniza el catálogo municipal oficial de Castilla y León",
    )
    parser.add_argument(
        "--directory-year",
        type=int,
        choices=sorted(INE_DIRECTORY_SOURCE_SPECS),
        default=max(INE_DIRECTORY_SOURCE_SPECS),
        help="Anualidad del directorio oficial",
    )
    parser.add_argument(
        "--directory-workbook",
        type=Path,
        help="XLSX local del directorio; por defecto se descarga del INE",
    )
    parser.add_argument(
        "--population-year",
        type=int,
        choices=sorted(INE_SOURCE_SPECS),
        default=max(INE_SOURCE_SPECS),
        help="Anualidad oficial de población",
    )
    parser.add_argument(
        "--population-archive",
        type=Path,
        help="ZIP local de población; por defecto se descarga del INE",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Aplica el plan; sin esta opción siempre se ejecuta en dry-run",
    )
    parser.add_argument(
        "--overwrite-existing",
        action="store_true",
        help="Permite sustituir población existente sin procedencia",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    directory_spec = INE_DIRECTORY_SOURCE_SPECS[args.directory_year]
    population_spec = INE_SOURCE_SPECS[args.population_year]
    try:
        directory_bytes = (
            read_local_workbook(args.directory_workbook)
            if args.directory_workbook is not None
            else download_directory_workbook(directory_spec)
        )
        directory_records, directory_provenance = (
            load_castilla_y_leon_directory(directory_bytes, directory_spec)
        )
        population_bytes = (
            read_local_archive(args.population_archive)
            if args.population_archive is not None
            else download_ine_archive(population_spec)
        )
        population_records, population_provenance = load_ine_population(
            population_bytes, population_spec
        )
        with SessionLocal() as db:
            plan = build_directory_sync_plan(
                db,
                directory_records,
                population_records,
                directory_provenance,
                population_provenance,
                overwrite_existing=args.overwrite_existing,
            )
            if args.apply and not plan.blocking_issues:
                apply_directory_sync(db, plan)
            applied = args.apply and not plan.blocking_issues
            summary = plan.summary(applied=applied)
        print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))
        return 3 if plan.blocking_issues else 0
    except (IneDirectoryError, InePopulationError) as exc:
        print(
            json.dumps(
                {"ok": False, "error": str(exc)},
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
        )
        return 2
    except SQLAlchemyError:
        print(
            json.dumps(
                {"ok": False, "error": "No se pudo completar la transacción"},
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
        )
        return 4


if __name__ == "__main__":
    raise SystemExit(main())
