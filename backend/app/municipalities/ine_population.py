"""Synchronize official INE municipal population figures.

The command is deliberately a dry-run unless ``--apply`` is supplied. It only
updates existing municipalities matched by their five-digit INE code; it never
creates municipalities or guesses matches from names.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import Request, urlopen
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.db.model_registry import register_all_models
from app.db.session import SessionLocal
from app.municipalities.models import Municipality

MAX_ARCHIVE_BYTES = 20 * 1024 * 1024
MAX_WORKBOOK_BYTES = 5 * 1024 * 1024
DOWNLOAD_TIMEOUT_SECONDS = 30.0
DOWNLOAD_CHUNK_BYTES = 64 * 1024
INE_CODE_PATTERN = re.compile(r"^[0-9]{5}$")


class InePopulationError(ValueError):
    """The official source or synchronization plan is not safe to apply."""


@dataclass(frozen=True)
class IneSourceSpec:
    year: int
    archive_url: str
    workbook_member: str
    title: str
    headers: tuple[str, ...]
    expected_sha256: str
    expected_rows: int


INE_SOURCE_SPECS = {
    2025: IneSourceSpec(
        year=2025,
        archive_url="https://www.ine.es/pob_xls/pobmun.zip",
        workbook_member="pobmun/pobmun25.xlsx",
        title=(
            "Cifras de población resultantes de la Revisión del Padrón "
            "municipal a 1 de enero de 2025"
        ),
        headers=(
            "CPRO",
            "PROVINCIA",
            "CMUN",
            "NOMBRE",
            "POB25",
            "HOMBRES",
            "MUJERES",
        ),
        expected_sha256=(
            "34d62a036d16dddac811e98b3edf341e3f9abe1fbf51f646c06f42947c5352da"
        ),
        expected_rows=8132,
    )
}


@dataclass(frozen=True)
class InePopulationRecord:
    ine_code: str
    province: str
    name: str
    population: int
    men: int
    women: int


@dataclass(frozen=True)
class PopulationProvenance:
    reference_year: int
    source_url: str
    source_sha256: str


@dataclass(frozen=True)
class SyncIssue:
    municipality_id: int
    ine_code: str | None
    municipality_name: str
    reason: str
    detail: str


@dataclass
class PopulationUpdate:
    municipality: Municipality
    record: InePopulationRecord
    density: float | None


@dataclass
class PopulationSyncPlan:
    provenance: PopulationProvenance
    source_rows: int
    database_rows: int
    matched_rows: int = 0
    unchanged_rows: int = 0
    updates: list[PopulationUpdate] = field(default_factory=list)
    missing_ine_codes: list[SyncIssue] = field(default_factory=list)
    invalid_ine_codes: list[SyncIssue] = field(default_factory=list)
    missing_source_codes: list[SyncIssue] = field(default_factory=list)
    name_mismatches: list[SyncIssue] = field(default_factory=list)
    conflicts: list[SyncIssue] = field(default_factory=list)

    def summary(self, *, applied: bool) -> dict[str, Any]:
        return {
            "ok": not self.conflicts,
            "mode": "apply" if applied else "dry-run",
            "source": {
                "year": self.provenance.reference_year,
                "url": self.provenance.source_url,
                "sha256": self.provenance.source_sha256,
                "rows": self.source_rows,
            },
            "database": {
                "rows": self.database_rows,
                "matched": self.matched_rows,
                "would_update": len(self.updates),
                "updated": len(self.updates) if applied else 0,
                "unchanged": self.unchanged_rows,
            },
            "issues": {
                "missing_ine_codes": [
                    asdict(issue) for issue in self.missing_ine_codes
                ],
                "invalid_ine_codes": [
                    asdict(issue) for issue in self.invalid_ine_codes
                ],
                "missing_source_codes": [
                    asdict(issue) for issue in self.missing_source_codes
                ],
                "name_mismatches": [asdict(issue) for issue in self.name_mismatches],
                "conflicts": [asdict(issue) for issue in self.conflicts],
            },
        }


def download_ine_archive(
    spec: IneSourceSpec,
    *,
    max_bytes: int = MAX_ARCHIVE_BYTES,
    timeout_seconds: float = DOWNLOAD_TIMEOUT_SECONDS,
) -> bytes:
    request = Request(
        spec.archive_url,
        headers={"User-Agent": "asistente-ayuntamientos-ine-sync/1.0"},
    )
    try:
        with urlopen(request, timeout=timeout_seconds) as response:
            final_url = urlsplit(response.geturl())
            expected_url = urlsplit(spec.archive_url)
            if (
                final_url.scheme != "https"
                or final_url.hostname != expected_url.hostname
            ):
                raise InePopulationError(
                    "La descarga del INE redirigió a un origen no permitido"
                )

            content_length = response.headers.get("Content-Length")
            if content_length is not None and int(content_length) > max_bytes:
                raise InePopulationError("El ZIP del INE supera el tamaño permitido")

            archive = bytearray()
            while chunk := response.read(DOWNLOAD_CHUNK_BYTES):
                archive.extend(chunk)
                if len(archive) > max_bytes:
                    raise InePopulationError(
                        "El ZIP del INE supera el tamaño permitido"
                    )
    except InePopulationError:
        raise
    except (HTTPError, URLError, OSError, ValueError) as exc:
        raise InePopulationError(f"No se pudo descargar el ZIP del INE: {exc}") from exc

    if not archive:
        raise InePopulationError("El ZIP descargado del INE está vacío")
    return bytes(archive)


def read_local_archive(
    path: Path,
    *,
    max_bytes: int = MAX_ARCHIVE_BYTES,
) -> bytes:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise InePopulationError(f"No se pudo leer el ZIP local: {exc}") from exc
    if size <= 0:
        raise InePopulationError("El ZIP local está vacío")
    if size > max_bytes:
        raise InePopulationError("El ZIP local supera el tamaño permitido")
    try:
        return path.read_bytes()
    except OSError as exc:
        raise InePopulationError(f"No se pudo leer el ZIP local: {exc}") from exc


def extract_workbook(
    archive_bytes: bytes,
    spec: IneSourceSpec,
    *,
    max_workbook_bytes: int = MAX_WORKBOOK_BYTES,
) -> tuple[bytes, str]:
    if len(archive_bytes) > MAX_ARCHIVE_BYTES:
        raise InePopulationError("El ZIP del INE supera el tamaño permitido")
    try:
        with ZipFile(io.BytesIO(archive_bytes)) as archive:
            try:
                member = archive.getinfo(spec.workbook_member)
            except KeyError as exc:
                raise InePopulationError(
                    f"El ZIP no contiene {spec.workbook_member}"
                ) from exc
            if member.flag_bits & 0x1:
                raise InePopulationError("El XLSX del INE no puede estar cifrado")
            if member.file_size <= 0 or member.file_size > max_workbook_bytes:
                raise InePopulationError("El XLSX del INE tiene un tamaño no permitido")
            workbook_bytes = archive.read(member)
    except InePopulationError:
        raise
    except (BadZipFile, OSError) as exc:
        raise InePopulationError(
            f"El archivo descargado no es un ZIP válido: {exc}"
        ) from exc

    if len(workbook_bytes) != member.file_size:
        raise InePopulationError("El XLSX extraído del INE está incompleto")
    workbook_sha256 = hashlib.sha256(workbook_bytes).hexdigest()
    if workbook_sha256 != spec.expected_sha256:
        raise InePopulationError(
            "El XLSX del INE no coincide con la versión revisada; "
            f"esperado {spec.expected_sha256}, recibido {workbook_sha256}"
        )
    return workbook_bytes, workbook_sha256


def parse_workbook(
    workbook_bytes: bytes,
    spec: IneSourceSpec,
) -> dict[str, InePopulationRecord]:
    try:
        workbook = load_workbook(
            io.BytesIO(workbook_bytes),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise InePopulationError(f"No se pudo abrir el XLSX del INE: {exc}") from exc

    try:
        if len(workbook.worksheets) != 1:
            raise InePopulationError("El XLSX del INE debe contener una sola hoja")
        rows = iter(workbook.worksheets[0].iter_rows(values_only=True))
        try:
            title_row = next(rows)
            header_row = next(rows)
        except StopIteration as exc:
            raise InePopulationError("El XLSX del INE no contiene cabeceras") from exc

        title = _required_text(title_row[0] if title_row else None, "título", 1)
        if title != spec.title:
            raise InePopulationError(
                f"Título inesperado en el XLSX del INE: {title!r}"
            )
        headers = tuple(
            _required_text(value, f"cabecera {index + 1}", 2)
            for index, value in enumerate(header_row[: len(spec.headers)])
        )
        if headers != spec.headers:
            raise InePopulationError(
                f"Cabeceras inesperadas en el XLSX del INE: {headers!r}"
            )

        records: dict[str, InePopulationRecord] = {}
        for row_number, row in enumerate(rows, start=3):
            values = tuple(row[: len(spec.headers)])
            if not values or all(value is None or value == "" for value in values):
                continue
            if len(values) != len(spec.headers):
                raise InePopulationError(
                    f"Fila {row_number}: se esperaban {len(spec.headers)} columnas"
                )

            province_code = _code_component(values[0], 2, "CPRO", row_number)
            province = _required_text(values[1], "PROVINCIA", row_number)
            municipality_code = _code_component(values[2], 3, "CMUN", row_number)
            name = _required_text(values[3], "NOMBRE", row_number)
            population = _non_negative_integer(values[4], "POB25", row_number)
            men = _non_negative_integer(values[5], "HOMBRES", row_number)
            women = _non_negative_integer(values[6], "MUJERES", row_number)
            if men + women != population:
                raise InePopulationError(
                    f"Fila {row_number}: HOMBRES + MUJERES no coincide con POB25"
                )

            ine_code = f"{province_code}{municipality_code}"
            if ine_code in records:
                raise InePopulationError(
                    f"Fila {row_number}: código INE duplicado {ine_code}"
                )
            records[ine_code] = InePopulationRecord(
                ine_code=ine_code,
                province=province,
                name=name,
                population=population,
                men=men,
                women=women,
            )
    finally:
        workbook.close()

    if len(records) != spec.expected_rows:
        raise InePopulationError(
            "Número inesperado de municipios en el XLSX del INE: "
            f"esperado {spec.expected_rows}, recibido {len(records)}"
        )
    return records


def load_ine_population(
    archive_bytes: bytes,
    spec: IneSourceSpec,
) -> tuple[dict[str, InePopulationRecord], PopulationProvenance]:
    workbook_bytes, workbook_sha256 = extract_workbook(archive_bytes, spec)
    records = parse_workbook(workbook_bytes, spec)
    return records, PopulationProvenance(
        reference_year=spec.year,
        source_url=spec.archive_url,
        source_sha256=workbook_sha256,
    )


def build_population_sync_plan(
    db: Session,
    records: dict[str, InePopulationRecord],
    provenance: PopulationProvenance,
    *,
    overwrite_existing: bool = False,
) -> PopulationSyncPlan:
    register_all_models()
    municipalities = list(
        db.scalars(select(Municipality).order_by(Municipality.id)).all()
    )
    plan = PopulationSyncPlan(
        provenance=provenance,
        source_rows=len(records),
        database_rows=len(municipalities),
    )

    for municipality in municipalities:
        ine_code = (municipality.ine_code or "").strip()
        if not ine_code:
            plan.missing_ine_codes.append(
                _issue(municipality, "missing_ine_code", "No consta código INE")
            )
            continue
        if not INE_CODE_PATTERN.fullmatch(ine_code):
            plan.invalid_ine_codes.append(
                _issue(
                    municipality,
                    "invalid_ine_code",
                    "El código INE debe tener cinco dígitos",
                )
            )
            continue

        record = records.get(ine_code)
        if record is None:
            plan.missing_source_codes.append(
                _issue(
                    municipality,
                    "missing_source_code",
                    "El código no aparece en la fuente INE seleccionada",
                )
            )
            continue

        plan.matched_rows += 1
        if _normalized_name(municipality.name) != _normalized_name(record.name):
            plan.name_mismatches.append(
                _issue(
                    municipality,
                    "name_mismatch",
                    f"INE: {record.name} ({record.province})",
                )
            )

        conflict = _population_conflict(
            municipality,
            record,
            provenance,
            overwrite_existing=overwrite_existing,
        )
        if conflict is not None:
            plan.conflicts.append(conflict)
            continue

        density = _density(record.population, municipality.surface_km2)
        desired = (
            record.population,
            provenance.reference_year,
            provenance.source_url,
            provenance.source_sha256,
            density,
        )
        current = (
            municipality.population,
            municipality.population_reference_year,
            municipality.population_source_url,
            municipality.population_source_sha256,
            municipality.density,
        )
        if current == desired:
            plan.unchanged_rows += 1
            continue
        plan.updates.append(
            PopulationUpdate(
                municipality=municipality,
                record=record,
                density=density,
            )
        )

    return plan


def apply_population_sync(db: Session, plan: PopulationSyncPlan) -> None:
    if plan.conflicts:
        raise InePopulationError(
            "La sincronización contiene conflictos; revisa el dry-run antes de aplicar"
        )
    try:
        for update in plan.updates:
            municipality = update.municipality
            municipality.population = update.record.population
            municipality.population_reference_year = plan.provenance.reference_year
            municipality.population_source_url = plan.provenance.source_url
            municipality.population_source_sha256 = plan.provenance.source_sha256
            municipality.density = update.density
        db.commit()
    except Exception:
        db.rollback()
        raise


def _population_conflict(
    municipality: Municipality,
    record: InePopulationRecord,
    provenance: PopulationProvenance,
    *,
    overwrite_existing: bool,
) -> SyncIssue | None:
    existing_provenance = (
        municipality.population_reference_year,
        municipality.population_source_url,
        municipality.population_source_sha256,
    )
    present_fields = sum(value is not None for value in existing_provenance)
    if present_fields not in (0, 3):
        return _issue(
            municipality,
            "partial_provenance",
            "La procedencia existente está incompleta",
        )

    if present_fields == 0:
        if (
            municipality.population is not None
            and municipality.population != record.population
            and not overwrite_existing
        ):
            return _issue(
                municipality,
                "unprovenanced_population",
                "La población existente "
                f"({municipality.population}) difiere de INE ({record.population})",
            )
        return None

    existing_year = municipality.population_reference_year
    existing_url = municipality.population_source_url
    existing_sha256 = municipality.population_source_sha256
    if existing_url != provenance.source_url:
        return _issue(
            municipality,
            "different_source",
            f"La procedencia existente es {existing_url}",
        )
    if existing_year is not None and existing_year > provenance.reference_year:
        return _issue(
            municipality,
            "newer_population",
            f"La población existente corresponde a {existing_year}",
        )
    if existing_year == provenance.reference_year:
        if existing_sha256 != provenance.source_sha256:
            return _issue(
                municipality,
                "different_source_revision",
                "La misma anualidad procede de otro fichero",
            )
        if municipality.population != record.population:
            return _issue(
                municipality,
                "inconsistent_official_population",
                "La cifra guardada no coincide con el fichero que figura como fuente",
            )
    return None


def _issue(
    municipality: Municipality,
    reason: str,
    detail: str,
) -> SyncIssue:
    return SyncIssue(
        municipality_id=municipality.id,
        ine_code=municipality.ine_code,
        municipality_name=municipality.name,
        reason=reason,
        detail=detail,
    )


def _density(population: int, surface_km2: float | None) -> float | None:
    if surface_km2 is None or surface_km2 <= 0:
        return None
    return population / surface_km2


def _normalized_name(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.casefold())
    without_marks = "".join(
        character
        for character in decomposed
        if not unicodedata.combining(character)
    )
    return " ".join(re.sub(r"[^a-z0-9]+", " ", without_marks).split())


def _required_text(value: object, field_name: str, row_number: int) -> str:
    text = str(value or "").strip()
    if not text:
        raise InePopulationError(f"Fila {row_number}: {field_name} está vacío")
    return text


def _code_component(
    value: object,
    width: int,
    field_name: str,
    row_number: int,
) -> str:
    if isinstance(value, bool):
        raise InePopulationError(f"Fila {row_number}: {field_name} no es válido")
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value or "").strip()
    if not text.isdigit() or len(text) > width:
        raise InePopulationError(f"Fila {row_number}: {field_name} no es válido")
    return text.zfill(width)


def _non_negative_integer(
    value: object,
    field_name: str,
    row_number: int,
) -> int:
    if isinstance(value, bool):
        raise InePopulationError(f"Fila {row_number}: {field_name} no es un entero")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float) and value.is_integer():
        parsed = int(value)
    else:
        text = str(value or "").strip()
        if not text.isdigit():
            raise InePopulationError(
                f"Fila {row_number}: {field_name} no es un entero"
            )
        parsed = int(text)
    if parsed < 0:
        raise InePopulationError(f"Fila {row_number}: {field_name} es negativo")
    return parsed


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Sincroniza cifras oficiales de población municipal del INE",
    )
    parser.add_argument(
        "--year",
        type=int,
        choices=sorted(INE_SOURCE_SPECS),
        default=max(INE_SOURCE_SPECS),
        help="Anualidad oficial que se importará",
    )
    parser.add_argument(
        "--archive",
        type=Path,
        help="ZIP local previamente descargado; por defecto se descarga del INE",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Aplica el plan; sin esta opción siempre se ejecuta en dry-run",
    )
    parser.add_argument(
        "--overwrite-existing",
        action="store_true",
        help="Permite sustituir cifras existentes que carecen de procedencia",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    spec = INE_SOURCE_SPECS[args.year]
    try:
        archive_bytes = (
            read_local_archive(args.archive)
            if args.archive is not None
            else download_ine_archive(spec)
        )
        records, provenance = load_ine_population(archive_bytes, spec)
        with SessionLocal() as db:
            plan = build_population_sync_plan(
                db,
                records,
                provenance,
                overwrite_existing=args.overwrite_existing,
            )
            if args.apply and not plan.conflicts:
                apply_population_sync(db, plan)
            summary = plan.summary(applied=args.apply and not plan.conflicts)
        print(json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True))
        return 3 if plan.conflicts else 0
    except InePopulationError as exc:
        print(
            json.dumps(
                {"ok": False, "error": str(exc)},
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
        )
        return 2
    except SQLAlchemyError:
        print(
            json.dumps(
                {"ok": False, "error": "No se pudo completar la transacción"},
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
        )
        return 4


if __name__ == "__main__":
    raise SystemExit(main())
